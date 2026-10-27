import os
import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def set_cell_margins(cell, top=100, bottom=100, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_apa_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="8" w:space="0" w:color="000000"/>
            <w:bottom w:val="single" w:sz="8" w:space="0" w:color="000000"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
            <w:insideH w:val="none"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def set_row_bottom_border(row, sz="6"):
    for cell in row.cells:
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = parse_xml(f'''
            <w:tcBorders {nsdecls("w")}>
                <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="000000"/>
            </w:tcBorders>
        ''')
        tcPr.append(tcBorders)

def build_perfect_apa_docx(file_path):
    doc = docx.Document(file_path)
    body = doc._body._element

    # 1. Habilitar w:updateFields para que Word actualice campos al abrir
    settings = doc.settings._element
    update_fields = parse_xml(r'<w:updateFields %s w:val="true"/>' % nsdecls('w'))
    settings.append(update_fields)

    # 2. Desempaquetar tablas flotantes de Quarto a formato nativo APA 7 de dos líneas
    outer_tables = list(doc.tables)
    for t_idx, outer_tbl in enumerate(outer_tables):
        if len(outer_tbl.rows) == 1 and len(outer_tbl.rows[0].cells) == 1:
            cell = outer_tbl.rows[0].cells[0]
            nested_tables = cell.tables
            if nested_tables:
                nested_tbl = nested_tables[0]
                caption_p = cell.paragraphs[0]
                raw_caption = caption_p.text.strip().replace('\xa0', ' ')

                if ':' in raw_caption:
                    parts = raw_caption.split(':', 1)
                    num_part = parts[0].strip().replace('Table', 'Tabla')
                    title_part = parts[1].strip()
                else:
                    num_part = f"Tabla {t_idx + 1}"
                    title_part = raw_caption

                p1_xml = f'''
                <w:p {nsdecls("w")}>
                    <w:pPr>
                        <w:jc w:val="left"/>
                        <w:spacing w:before="240" w:after="40" w:line="240" w:lineRule="auto"/>
                        <w:ind w:left="0" w:firstLine="0"/>
                    </w:pPr>
                    <w:r>
                        <w:rPr>
                            <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:cs="Times New Roman"/>
                            <w:b/>
                            <w:bCs/>
                            <w:sz w:val="24"/>
                            <w:szCs w:val="24"/>
                            <w:color w:val="000000"/>
                        </w:rPr>
                        <w:t>{num_part}</w:t>
                    </w:r>
                </w:p>
                '''

                p2_xml = f'''
                <w:p {nsdecls("w")}>
                    <w:pPr>
                        <w:jc w:val="left"/>
                        <w:spacing w:before="0" w:after="120" w:line="240" w:lineRule="auto"/>
                        <w:ind w:left="0" w:firstLine="0"/>
                    </w:pPr>
                    <w:r>
                        <w:rPr>
                            <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:cs="Times New Roman"/>
                            <w:i/>
                            <w:iCs/>
                            <w:sz w:val="24"/>
                            <w:szCs w:val="24"/>
                            <w:color w:val="000000"/>
                        </w:rPr>
                        <w:t>{title_part}</w:t>
                    </w:r>
                </w:p>
                '''

                p1 = parse_xml(p1_xml)
                p2 = parse_xml(p2_xml)
                nested_xml_elem = nested_tbl._tbl

                parent = outer_tbl._tbl.getparent()
                idx = parent.index(outer_tbl._tbl)

                parent.insert(idx, p1)
                parent.insert(idx + 1, p2)
                parent.insert(idx + 2, nested_xml_elem)
                parent.remove(outer_tbl._tbl)

    # 3. Formateo y estilo de todas las tablas restantes (10 pt exactos, Times New Roman, bordes APA 7)
    for t in doc.tables:
        set_apa_table_borders(t)
        if len(t.rows) > 0:
            set_row_bottom_border(t.rows[0], sz="6")

        for r_idx, row in enumerate(t.rows):
            for cell in row.cells:
                set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
                for p in cell.paragraphs:
                    p.paragraph_format.first_line_indent = Inches(0)
                    p.paragraph_format.left_indent = Inches(0)
                    p.paragraph_format.line_spacing = 1.15
                    p.paragraph_format.space_before = Pt(1)
                    p.paragraph_format.space_after = Pt(1)

                    for r in p.runs:
                        r.font.name = "Times New Roman"
                        r.font.size = Pt(10)
                        r.font.color.rgb = RGBColor(0, 0, 0)
                        if r_idx == 0:
                            r.font.bold = True

    # 4. Formatear párrafos del documento (Justificado, 12pt Times New Roman, sangría 0.5")
    in_ref_section = False
    for p in doc.paragraphs:
        txt = p.text.strip()
        s_name = p.style.name

        if s_name.startswith('Heading') and 'referencia' in txt.lower():
            in_ref_section = True
            continue
        elif s_name.startswith('Heading') and in_ref_section:
            in_ref_section = False

        if in_ref_section and txt:
            pPr = p._p.get_or_add_pPr()
            numPr = pPr.find(qn('w:numPr'))
            if numPr is not None:
                pPr.remove(numPr)

            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Inches(0.5)
            p.paragraph_format.first_line_indent = Inches(-0.5)
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(6)

            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(12)
                r.font.color.rgb = RGBColor(0, 0, 0)
            continue

        if s_name in ['Normal', 'Body Text', 'First Paragraph']:
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.first_line_indent = Inches(0.5)

            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(12)
                r.font.color.rgb = RGBColor(0, 0, 0)

        elif s_name.startswith('Heading'):
            p.paragraph_format.first_line_indent = Inches(0)
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.color.rgb = RGBColor(0, 0, 0)

    # 5. Configurar y limpiar la Tabla de Contenidos nativa de Word
    # Mantener el campo nativo de Word sin duplicar manualmente párrafos
    sdt_elements = body.xpath('.//w:sdt')
    if sdt_elements:
        sdt = sdt_elements[0]
        sdtContent = sdt.find(qn('w:sdtContent'))
        if sdtContent is not None:
            # Vaciar cualquier residuo estático
            for child in list(sdtContent):
                sdtContent.remove(child)

            # Título nativo de TOC en español
            title_p = parse_xml(f'''
            <w:p {nsdecls("w")}>
                <w:pPr>
                    <w:pStyle w:val="TOCHeading"/>
                    <w:jc w:val="center"/>
                    <w:spacing w:before="120" w:after="240"/>
                </w:pPr>
                <w:r>
                    <w:rPr>
                        <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:cs="Times New Roman"/>
                        <w:b/>
                        <w:sz w:val="24"/>
                        <w:szCs w:val="24"/>
                        <w:color w:val="000000"/>
                    </w:rPr>
                    <w:t>Índice de Contenido</w:t>
                </w:r>
            </w:p>
            ''')
            sdtContent.append(title_p)

            # Instrucción de campo nativa de Word para TOC con hipervínculos automáticos y niveles 1-3
            fld_p = parse_xml(f'''
            <w:p {nsdecls("w")}>
                <w:r>
                    <w:fldChar w:fldCharType="begin" w:dirty="true"/>
                    <w:instrText xml:space="preserve"> TOC \\o "1-3" \\h \\z \\u </w:instrText>
                    <w:fldChar w:fldCharType="separate"/>
                    <w:fldChar w:fldCharType="end"/>
                </w:r>
            </w:p>
            ''')
            sdtContent.append(fld_p)

            # Salto de página formal tras el índice
            page_break = parse_xml(f'''
            <w:p {nsdecls("w")}>
                <w:r>
                    <w:br w:type="page"/>
                </w:r>
            </w:p>
            ''')
            sdtContent.append(page_break)

    doc.save(file_path)
    print(f"Successfully processed clean document: {file_path}")

if __name__ == "__main__":
    target = r"d:\08_Ciclo\03_Practicas_Preprofesionales\docs\vaults\v04-expediente-proyecto-capuli\plantillas\plantilla_metodologia_giz_apa.docx"
    build_perfect_apa_docx(target)
