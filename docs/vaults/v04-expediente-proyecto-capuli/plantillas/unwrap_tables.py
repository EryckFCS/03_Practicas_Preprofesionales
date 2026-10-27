import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

doc_path = r"d:\08_Ciclo\03_Practicas_Preprofesionales\docs\vaults\v04-expediente-proyecto-capuli\plantillas\plantilla_metodologia_giz_apa.docx"
doc = docx.Document(doc_path)
body = doc._body._element

# Buscar todas las tablas de primer nivel
outer_tables = list(doc.tables)
print(f"Inspecting {len(outer_tables)} outer tables...")

for t_idx, outer_tbl in enumerate(outer_tables):
    cell = outer_tbl.rows[0].cells[0]
    nested_tables = cell.tables
    if not nested_tables:
        continue
    
    nested_tbl = nested_tables[0]
    
    # Extraer el texto de la leyenda
    caption_p = cell.paragraphs[0]
    raw_caption = caption_p.text.strip().replace('\xa0', ' ')
    print(f"Table {t_idx} raw caption: {raw_caption}")
    
    # Separar número y título según APA 7 (Ej: 'Table 1: Identificación general...')
    # Línea 1: Tabla 1 (Negrita, sin cursiva)
    # Línea 2: Identificación general... (Cursiva, sin negrita)
    if ':' in raw_caption:
        parts = raw_caption.split(':', 1)
        num_part = parts[0].strip().replace('Table', 'Tabla')
        title_part = parts[1].strip()
    else:
        num_part = f"Tabla {t_idx + 1}"
        title_part = raw_caption

    # Crear los dos párrafos reglamentarios APA 7 para colocar ANTES de la tabla
    p1_xml = f'''
    <w:p {nsdecls("w")}>
        <w:pPr>
            <w:pStyle w:val="TableCaption"/>
            <w:jc w:val="left"/>
            <w:spacing w:before="240" w:after="40" w:line="240" w:lineRule="auto"/>
            <w:ind w:left="0" w:firstLine="0"/>
        </w:pPr>
        <w:r>
            <w:rPr>
                <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:cs="Times New Roman"/>
                <w:b/>
                <w:bCs/>
                <w:sz w:val="24"/>
                <w:szCs w:val="24"/>
                <w:color w:val="000000"/>
            </w:rPr>
            <w:t>{num_part}</w:t>
        </w:r>
    </w:p>
    '''
    
    p2_xml = f'''
    <w:p {nsdecls("w")}>
        <w:pPr>
            <w:pStyle w:val="TableCaption"/>
            <w:jc w:val="left"/>
            <w:spacing w:before="0" w:after="120" w:line="240" w:lineRule="auto"/>
            <w:ind w:left="0" w:firstLine="0"/>
        </w:pPr>
        <w:r>
            <w:rPr>
                <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:cs="Times New Roman"/>
                <w:i/>
                <w:iCs/>
                <w:sz w:val="24"/>
                <w:szCs w:val="24"/>
                <w:color w:val="000000"/>
            </w:rPr>
            <w:t>{title_part}</w:t>
        </w:r>
    </w:p>
    '''
    
    p1 = parse_xml(p1_xml)
    p2 = parse_xml(p2_xml)
    
    # Extraer el elemento XML de la tabla anidada real
    nested_xml_elem = nested_tbl._tbl
    
    # Ubicar outer_tbl en body
    parent = outer_tbl._tbl.getparent()
    idx = parent.index(outer_tbl._tbl)
    
    # Insertar p1, p2 y la tabla real en el body, reemplazando a outer_tbl
    parent.insert(idx, p1)
    parent.insert(idx + 1, p2)
    parent.insert(idx + 2, nested_xml_elem)
    parent.remove(outer_tbl._tbl)

doc.save(doc_path)
print("Unwrapped all tables successfully into native APA 7 two-line format!")
